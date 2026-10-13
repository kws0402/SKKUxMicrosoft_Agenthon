from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

import openpyxl

from scripts.download_excel import download_url
from scripts.fleet_sync import (
    FleetDataError,
    SHEET_HEADERS,
    build_payload,
    validate_payload,
    write_javascript_atomic,
    write_json_atomic,
)


def row_for(sheet: str, **overrides: object) -> list[object]:
    defaults: dict[str, object] = {
        "vessel_name": "TEST SPIRIT",
        "vessel_id": "V-001",
        "mmsi": "440000001",
        "imo_number": "9000001",
        "call_sign": "TEST1",
        "vessel_type": "PCTC",
        "flag_country": "KR",
        "capacity": 100,
        "loaded_vehicle": 10,
        "remaining_capacity": 90,
        "operating_status": "배선 가능",
        "home_port_code": "P1",
        "max_speed_knots": 20,
        "effective_speed_cap_knots": 18,
        "environmental_speed_restriction": True,
        "environmental_rule": "TEST",
        "design_draft_m": 9.5,
        "loa_m": 200,
        "year_built": 2024,
        "fleet_priority": "HIGH",
        "far_east_return_required": True,
        "availability_note": "Available",
        "data_source": "TEST",
        "last_updated": "2026-10-13T00:00:00",
        "service_speed_knots": 17,
        "current_max_speed_knots": 18,
        "vessel_draft": 9.5,
        "status": "AVAILABLE",
        "available_capacity": 84,
        "maintenance_start_date": None,
        "maintenance_end_date": None,
        "status_before_reservation": None,
        "deck1_large_capacity": 20,
        "deck1_large_available": 18,
        "deck2_medium_capacity": 20,
        "deck2_medium_available": 18,
        "deck3_small_capacity": 20,
        "deck3_small_available": 18,
        "voyage_id": "VOY-001",
        "reported_port_code": "P1",
        "latitude": 35.0,
        "longitude": 129.0,
        "speed_knots": 0.1,
        "course_deg": 0,
        "heading_deg": 0,
        "nav_status": "Moored",
        "position_source": "TEST",
        "port_name": "Port One",
        "port_code": "P1",
        "country": "KR",
        "max_draft_m": 15,
        "max_loa_m": 300,
        "region": "Northeast Asia",
        "terminal_type": "RoRo",
        "korea_major_port": True,
        "large_vessel_allowed": True,
        "port_restriction_note": "None",
        "previous_port_code": "P1",
        "current_port_code": "P1",
        "next_port_code": "P2",
        "final_destination_port_code": "P2",
        "route_plan": "P1 > P2",
        "voyage_purpose": "OWN_CARGO",
        "etd": "2026-10-13T00:00:00",
        "eta": "2026-10-14T00:00:00",
        "voyage_status": "출항 예정",
        "planned_speed_knots": 17,
        "environmental_speed_cap_knots": 18,
        "route_distance_nm": 100,
        "remaining_distance_nm": 100,
        "month_end_departure": False,
        "available_from_date": "2026-10-14T00:00:00",
        "available_port_code": "P2",
        "route_change_allowed": False,
        "order_id": "ORD-001",
        "customer": "Customer",
        "origin_port_code": "P1",
        "destination_port_code": "P2",
        "vehicle_type": "Mixed",
        "quantity": 6,
        "requested_departure": "2026-10-13",
        "loading_deadline": "2026-10-13",
        "delivery_deadline": "2026-10-14",
        "priority": "HIGH",
        "month_end_loading_importance": "NORMAL",
        "reporting_month": "2026-10-01",
        "assigned_vessel_id": "V-001",
        "assigned_voyage_id": "VOY-001",
        "order_status": "ASSIGNED",
        "assignment_reason": "Test",
        "created_at": "2026-10-01",
        "requested_loading_date": "2026-10-13",
        "latest_loading_date": "2026-10-13",
        "confirmation_status": "CONFIRMED",
        "confirmed_by": "Tester",
        "confirmed_at": "2026-10-12",
        "reservation_type": "TEST",
        "large_vehicle_qty": 2,
        "medium_vehicle_qty": 2,
        "small_vehicle_qty": 2,
        "charter_id": "CH-001",
        "charter_status": "AVAILABLE",
        "charter_company_id": "CO-001",
        "charter_start_date": None,
        "charter_end_date": None,
        "contact_id": "CONTACT-001",
        "note": "Test",
        "company_id": "CO-001",
        "company_name": "Carrier",
        "department": "Sales",
        "contact_name": "Contact",
        "phone": "000",
        "email": "test@example.invalid",
    }
    defaults.update(overrides)
    return [defaults.get(header) for header in SHEET_HEADERS[sheet]]


def make_workbook(path: Path) -> None:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for sheet_name, headers in SHEET_HEADERS.items():
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(headers)
        if sheet_name == "PORT":
            sheet.append(row_for(sheet_name, port_code="P1", port_name="Port One", latitude=35.0, longitude=129.0))
            sheet.append(row_for(sheet_name, port_code="P2", port_name="Port Two", latitude=1.2, longitude=103.8, country="SG"))
        elif sheet_name == "CHARTER":
            sheet.append(row_for(sheet_name, charter_status="CHARTERED_OUT"))
        else:
            sheet.append(row_for(sheet_name))
    workbook.save(path)


class FleetSyncTests(unittest.TestCase):
    def test_download_url_preserves_share_token_and_adds_download(self) -> None:
        result = download_url("https://example.test/file?e=token")
        self.assertEqual(result, "https://example.test/file?e=token&download=1")

    def test_build_payload_joins_and_recalculates_capacity(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workbook = Path(directory) / "source.xlsx"
            make_workbook(workbook)
            payload = build_payload(workbook, generated_at="2026-10-13T00:00:00Z")
        self.assertEqual(payload["meta"]["counts"]["VESSEL"], 1)
        self.assertEqual(payload["meta"]["filename"], "fleet-source.xlsx")
        vessel = payload["vessel"][0]
        self.assertEqual(vessel["available_capacity"], 84)
        self.assertEqual(
            [vessel["deck1_large_available"], vessel["deck2_medium_available"], vessel["deck3_small_available"]],
            [18, 18, 18],
        )
        self.assertEqual(payload["voyage"][0]["route_plan"], "P1 > P2")
        self.assertEqual(payload["charter_status"][0]["contact"]["company_name"], "Carrier")
        validate_payload(payload)

    def test_generated_at_is_stable_when_content_does_not_change(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workbook = Path(directory) / "source.xlsx"
            make_workbook(workbook)
            first = build_payload(workbook, generated_at="2026-10-13T00:00:00Z")
            second = build_payload(
                workbook,
                previous_payload=first,
                generated_at="2026-10-13T01:00:00Z",
            )
        self.assertEqual(second["meta"]["generated_at"], "2026-10-13T00:00:00Z")
        self.assertEqual(second["meta"]["content_hash"], first["meta"]["content_hash"])

    def test_header_change_fails_before_output_is_replaced(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook_path = root / "source.xlsx"
            output_path = root / "fleet.json"
            make_workbook(workbook_path)
            output_path.write_text('{"preserved": true}\n', encoding="utf-8")
            workbook = openpyxl.load_workbook(workbook_path)
            workbook["VESSEL"]["A1"] = "renamed_column"
            workbook.save(workbook_path)
            with self.assertRaises(FleetDataError):
                payload = build_payload(workbook_path)
                write_json_atomic(payload, output_path)
            self.assertEqual(json.loads(output_path.read_text(encoding="utf-8")), {"preserved": True})

    def test_invalid_hash_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workbook = Path(directory) / "source.xlsx"
            make_workbook(workbook)
            payload = build_payload(workbook)
        payload["meta"]["content_hash"] = "0" * 64
        with self.assertRaises(FleetDataError):
            validate_payload(payload)

    def test_browser_javascript_wraps_the_same_payload(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook = root / "source.xlsx"
            output = root / "fleet-data.js"
            make_workbook(workbook)
            payload = build_payload(workbook, generated_at="2026-10-13T00:00:00Z")
            write_javascript_atomic(payload, output)
            script = output.read_text(encoding="utf-8")
        self.assertTrue(script.startswith("window.FLEET_DATA="))
        self.assertTrue(script.endswith(";\n"))
        self.assertEqual(json.loads(script.removeprefix("window.FLEET_DATA=")[:-2]), payload)

    def test_browser_distinguishes_operating_and_planned_charter_status(self) -> None:
        index = (Path(__file__).parents[1] / "index.html").read_text(encoding="utf-8")
        self.assertIn("v.status==='CHARTER_PLANNED'", index)
        self.assertIn("plannedStatusBadge(v)", index)
        self.assertIn("${esc(v.operating_status)}${plannedStatusText(v)}", index)


if __name__ == "__main__":
    unittest.main()
