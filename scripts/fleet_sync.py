from __future__ import annotations

import hashlib
import json
import math
from collections import defaultdict
from datetime import date, datetime, time, timezone
from pathlib import Path
from typing import Any, Iterable

import openpyxl


SHEET_HEADERS: dict[str, tuple[str, ...]] = {
    "VESSEL": (
        "vessel_name", "vessel_id", "mmsi", "imo_number", "call_sign",
        "vessel_type", "flag_country", "capacity", "loaded_vehicle",
        "remaining_capacity", "operating_status", "home_port_code",
        "max_speed_knots", "effective_speed_cap_knots",
        "environmental_speed_restriction", "environmental_rule",
        "design_draft_m", "loa_m", "year_built", "fleet_priority",
        "far_east_return_required", "availability_note", "data_source",
        "last_updated", "service_speed_knots", "current_max_speed_knots",
        "vessel_draft", "status", "available_capacity",
        "maintenance_start_date", "maintenance_end_date",
        "status_before_reservation", "deck1_large_capacity",
        "deck1_large_available", "deck2_medium_capacity",
        "deck2_medium_available", "deck3_small_capacity",
        "deck3_small_available",
    ),
    "VESSEL_LIVE": (
        "vessel_id", "vessel_name", "mmsi", "voyage_id",
        "reported_port_code", "latitude", "longitude", "speed_knots",
        "course_deg", "heading_deg", "nav_status", "position_source",
        "last_updated",
    ),
    "PORT": (
        "port_name", "port_code", "country", "latitude", "longitude",
        "max_draft_m", "max_loa_m", "region", "terminal_type",
        "korea_major_port", "data_source", "last_updated",
        "large_vessel_allowed", "port_restriction_note",
    ),
    "VOYAGE": (
        "voyage_id", "vessel_id", "vessel_name", "previous_port_code",
        "current_port_code", "next_port_code", "final_destination_port_code",
        "route_plan", "voyage_purpose", "etd", "eta", "voyage_status",
        "planned_speed_knots", "environmental_speed_cap_knots",
        "route_distance_nm", "remaining_distance_nm", "month_end_departure",
        "fleet_priority", "far_east_return_required", "data_source",
        "last_updated", "available_from_date", "available_port_code",
        "route_change_allowed",
    ),
    "CARGO": (
        "order_id", "customer", "origin_port_code", "destination_port_code",
        "vehicle_type", "quantity", "requested_departure", "loading_deadline",
        "delivery_deadline", "priority", "month_end_loading_importance",
        "reporting_month", "assigned_vessel_id", "assigned_voyage_id",
        "order_status", "assignment_reason", "data_source", "created_at",
        "requested_loading_date", "latest_loading_date", "confirmation_status",
        "confirmed_by", "confirmed_at", "reservation_type",
        "large_vehicle_qty", "medium_vehicle_qty", "small_vehicle_qty",
    ),
    "CHARTER": (
        "charter_id", "vessel_id", "charter_status", "charter_company_id",
        "charter_start_date", "charter_end_date", "contact_id", "note",
        "data_source",
    ),
    "CHARTER_CONTACT": (
        "contact_id", "company_id", "company_name", "department",
        "contact_name", "phone", "email", "region", "note", "data_source",
    ),
}

PUBLIC_EXCEL_URL = (
    "https://m365cpi33505044-my.sharepoint.com/:x:/g/personal/"
    "billiev_m365cpi33505044_onmicrosoft_com/"
    "IQDQxuRGYrzaQIaIuS0QZuoeAVRBVdIJryX1OePSh9sTF2Q?e=vWhSQO"
)


class FleetDataError(ValueError):
    """Raised when the workbook or generated payload is inconsistent."""


def _json_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        raise FleetDataError("Workbook contains a non-finite number")
    return value


def _rows(worksheet: Any, expected_headers: tuple[str, ...]) -> list[dict[str, Any]]:
    iterator = worksheet.iter_rows(values_only=True)
    try:
        actual = tuple(next(iterator))
    except StopIteration as exc:
        raise FleetDataError(f"Sheet {worksheet.title} is empty") from exc

    if actual != expected_headers:
        missing = [name for name in expected_headers if name not in actual]
        extra = [name for name in actual if name not in expected_headers]
        raise FleetDataError(
            f"Sheet {worksheet.title} header mismatch; missing={missing}, extra={extra}"
        )

    records: list[dict[str, Any]] = []
    for row_number, values in enumerate(iterator, start=2):
        if not any(value is not None for value in values):
            continue
        if len(values) != len(expected_headers):
            raise FleetDataError(
                f"Sheet {worksheet.title} row {row_number} has {len(values)} cells; "
                f"expected {len(expected_headers)}"
            )
        records.append(
            {key: _json_value(value) for key, value in zip(expected_headers, values)}
        )
    return records


def _index(records: Iterable[dict[str, Any]], key: str, label: str) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for record in records:
        value = record.get(key)
        if not isinstance(value, str) or not value.strip():
            raise FleetDataError(f"{label} has a missing {key}")
        if value in result:
            raise FleetDataError(f"{label} has duplicate {key}: {value}")
        result[value] = record
    return result


def _require_number(record: dict[str, Any], field: str, label: str) -> float | int:
    value = record.get(field)
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise FleetDataError(f"{label} requires numeric {field}; got {value!r}")
    return value


def _route(voyage: dict[str, Any], ports: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    codes = [code.strip() for code in str(voyage.get("route_plan") or "").split(">")]
    route: list[dict[str, Any]] = []
    for code in filter(None, codes):
        port = ports.get(code)
        if port is None:
            raise FleetDataError(f"Voyage {voyage['voyage_id']} references unknown port {code}")
        route.append(
            {
                "port_code": code,
                "port_name": port["port_name"],
                "country": port["country"],
                "latitude": port["latitude"],
                "longitude": port["longitude"],
            }
        )
    if not route:
        raise FleetDataError(f"Voyage {voyage['voyage_id']} has an empty route_plan")
    return route


DATA_KEYS = (
    "vessel",
    "cargo",
    "port",
    "voyage",
    "vessel_live",
    "maintenance_schedule",
    "charter_status",
    "fleet_allocation",
    "port_compatibility",
)


def _content_hash(payload: dict[str, Any]) -> str:
    canonical = {key: payload[key] for key in DATA_KEYS}
    return hashlib.sha256(
        json.dumps(canonical, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()


def validate_payload(payload: dict[str, Any]) -> None:
    required_top = {*DATA_KEYS, "meta"}
    missing = sorted(required_top - payload.keys())
    if missing:
        raise FleetDataError(f"Generated JSON is missing keys: {missing}")

    expected_counts = {key.upper(): len(payload[key]) for key in DATA_KEYS}
    for key, value in expected_counts.items():
        if payload["meta"].get("counts", {}).get(key) != value:
            raise FleetDataError(f"meta.counts.{key} does not match the payload")

    vessel_ids = [vessel.get("vessel_id") for vessel in payload["vessel"]]
    if len(vessel_ids) != len(set(vessel_ids)):
        raise FleetDataError("Generated JSON contains duplicate vessel_id values")

    live_by_vessel = _index(payload["vessel_live"], "vessel_id", "vessel_live")
    voyage_by_vessel = _index(payload["voyage"], "vessel_id", "voyage")
    if set(vessel_ids) != set(live_by_vessel) or set(vessel_ids) != set(voyage_by_vessel):
        raise FleetDataError("Generated JSON does not have one live position and voyage per vessel")

    for vessel in payload["vessel"]:
        live = live_by_vessel[vessel["vessel_id"]]
        latitude = live.get("latitude")
        longitude = live.get("longitude")
        if not isinstance(latitude, (int, float)) or not -90 <= latitude <= 90:
            raise FleetDataError(f"Vessel {vessel['vessel_id']} has invalid latitude")
        if not isinstance(longitude, (int, float)) or not -180 <= longitude <= 180:
            raise FleetDataError(f"Vessel {vessel['vessel_id']} has invalid longitude")
        for capacity_key, available_key in (
            ("deck1_large_capacity", "deck1_large_available"),
            ("deck2_medium_capacity", "deck2_medium_available"),
            ("deck3_small_capacity", "deck3_small_available"),
        ):
            capacity = _require_number(vessel, capacity_key, f"Vessel {vessel['vessel_id']}")
            available = _require_number(vessel, available_key, f"Vessel {vessel['vessel_id']}")
            if capacity < 0 or available < 0 or available > capacity:
                raise FleetDataError(f"Vessel {vessel['vessel_id']} has invalid deck capacity")

    if _content_hash(payload) != payload["meta"].get("content_hash"):
        raise FleetDataError("contentHash does not match generated data")


def build_payload(
    workbook_path: str | Path,
    *,
    previous_payload: dict[str, Any] | None = None,
    source_url: str = PUBLIC_EXCEL_URL,
    generated_at: str | None = None,
) -> dict[str, Any]:
    path = Path(workbook_path)
    if not path.is_file():
        raise FleetDataError(f"Workbook does not exist: {path}")

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    missing_sheets = sorted(set(SHEET_HEADERS) - set(workbook.sheetnames))
    extra_sheets = sorted(set(workbook.sheetnames) - set(SHEET_HEADERS))
    if missing_sheets or extra_sheets:
        raise FleetDataError(
            f"Workbook sheet mismatch; missing={missing_sheets}, extra={extra_sheets}"
        )

    raw = {
        name: _rows(workbook[name], headers)
        for name, headers in SHEET_HEADERS.items()
    }
    workbook.close()

    vessels_by_id = _index(raw["VESSEL"], "vessel_id", "VESSEL")
    live_by_vessel = _index(raw["VESSEL_LIVE"], "vessel_id", "VESSEL_LIVE")
    ports_by_code = _index(raw["PORT"], "port_code", "PORT")
    voyages_by_id = _index(raw["VOYAGE"], "voyage_id", "VOYAGE")
    voyages_by_vessel = _index(raw["VOYAGE"], "vessel_id", "VOYAGE")
    charters_by_vessel = _index(raw["CHARTER"], "vessel_id", "CHARTER")
    contacts_by_id = _index(raw["CHARTER_CONTACT"], "contact_id", "CHARTER_CONTACT")

    for label, index in (
        ("VESSEL_LIVE", live_by_vessel),
        ("VOYAGE", voyages_by_vessel),
        ("CHARTER", charters_by_vessel),
    ):
        missing = sorted(set(vessels_by_id) - set(index))
        extra = sorted(set(index) - set(vessels_by_id))
        if missing or extra:
            raise FleetDataError(f"{label} vessel coverage mismatch; missing={missing}, extra={extra}")

    cargo_totals: dict[str, dict[str, float | int]] = defaultdict(
        lambda: {"quantity": 0, "large_vehicle_qty": 0, "medium_vehicle_qty": 0, "small_vehicle_qty": 0}
    )
    for order in raw["CARGO"]:
        vessel_id = order.get("assigned_vessel_id")
        voyage_id = order.get("assigned_voyage_id")
        if vessel_id is not None and vessel_id not in vessels_by_id:
            raise FleetDataError(f"Cargo {order['order_id']} references unknown vessel {vessel_id}")
        if voyage_id is not None and voyage_id not in voyages_by_id:
            raise FleetDataError(f"Cargo {order['order_id']} references unknown voyage {voyage_id}")
        if vessel_id and order.get("order_status") == "ASSIGNED":
            for field in cargo_totals[vessel_id]:
                cargo_totals[vessel_id][field] += _require_number(order, field, f"Cargo {order['order_id']}")

    vessels: list[dict[str, Any]] = []
    for vessel_id, original in vessels_by_id.items():
        vessel = dict(original)
        live = live_by_vessel[vessel_id]
        voyage = voyages_by_vessel[vessel_id]
        if live.get("voyage_id") != voyage.get("voyage_id"):
            raise FleetDataError(f"Live and voyage IDs disagree for {vessel_id}")

        totals = cargo_totals[vessel_id]
        calculated_available = _require_number(original, "remaining_capacity", f"Vessel {vessel_id}") - totals["quantity"]
        cached_available = original.get("available_capacity")
        if cached_available is not None and cached_available != calculated_available:
            raise FleetDataError(
                f"Vessel {vessel_id} cached available_capacity={cached_available} "
                f"but cargo calculation={calculated_available}"
            )
        vessel["available_capacity"] = calculated_available

        deck_specs = (
            ("deck1_large_capacity", "deck1_large_available", "large_vehicle_qty"),
            ("deck2_medium_capacity", "deck2_medium_available", "medium_vehicle_qty"),
            ("deck3_small_capacity", "deck3_small_available", "small_vehicle_qty"),
        )
        for capacity_key, available_key, cargo_key in deck_specs:
            capacity = _require_number(original, capacity_key, f"Vessel {vessel_id}")
            available = capacity - totals[cargo_key]
            cached = original.get(available_key)
            if cached is not None and cached != available:
                raise FleetDataError(
                    f"Vessel {vessel_id} cached {available_key}={cached} but cargo calculation={available}"
                )
            if available < 0:
                raise FleetDataError(f"Vessel {vessel_id} {capacity_key} is over capacity")
            vessel[available_key] = available
        vessels.append(vessel)

    maintenance_schedule: list[dict[str, Any]] = []
    for sequence, vessel in enumerate((v for v in vessels if v.get("maintenance_start_date") or v.get("maintenance_end_date")), start=1):
        voyage = voyages_by_vessel[vessel["vessel_id"]]
        maintenance_schedule.append(
            {
                "maintenance_id": f"MNT-2026-{sequence:03d}",
                "vessel_id": vessel["vessel_id"],
                "maintenance_port_code": voyage.get("current_port_code") or vessel.get("home_port_code"),
                "maintenance_type": "SCHEDULED_MAINTENANCE",
                "start_datetime": vessel.get("maintenance_start_date"),
                "end_datetime": vessel.get("maintenance_end_date"),
                "status": "IN_PROGRESS" if vessel.get("status") == "MAINTENANCE" else "SCHEDULED",
                "blocking_flag": vessel.get("status") == "MAINTENANCE",
                "planner_note": vessel.get("availability_note"),
                "data_source": vessel.get("data_source"),
                "last_updated": vessel.get("last_updated"),
            }
        )

    charter_status: list[dict[str, Any]] = []
    for charter in raw["CHARTER"]:
        if charter.get("charter_status") not in {"CHARTERED_OUT", "CHARTER_PLANNED"}:
            continue
        contact_id = charter.get("contact_id")
        contact = contacts_by_id.get(contact_id) if contact_id else None
        if contact_id and contact is None:
            raise FleetDataError(f"Charter {charter['charter_id']} references unknown contact {contact_id}")
        charter_status.append(
            {
                "charter_id": charter["charter_id"],
                "vessel_id": charter["vessel_id"],
                "charterer_name": contact.get("company_name") if contact else charter.get("charter_company_id"),
                "charter_type": "TIME_CHARTER_OUT",
                "start_datetime": charter.get("charter_start_date"),
                "end_datetime": charter.get("charter_end_date"),
                "charter_status": "ACTIVE" if charter.get("charter_status") == "CHARTERED_OUT" else "PLANNED",
                "own_cargo_allowed": False,
                "contract_reference": charter["charter_id"],
                "planner_note": charter.get("note"),
                "data_source": charter.get("data_source"),
                "last_updated": vessels_by_id[charter["vessel_id"]].get("last_updated"),
                "contact": contact,
            }
        )

    fleet_allocation: list[dict[str, Any]] = []
    for vessel in vessels:
        voyage = voyages_by_vessel[vessel["vessel_id"]]
        fleet_allocation.append(
            {
                "allocation_id": f"ALC-{vessel['vessel_id']}",
                "allocation_month": str(voyage.get("etd") or "")[:7] + "-01",
                "vessel_id": vessel["vessel_id"],
                "allocated_team": "FAR_EAST_TEAM" if vessel.get("far_east_return_required") else "GLOBAL_TEAM",
                "allocation_capacity": vessel.get("available_capacity"),
                "fleet_priority": vessel.get("fleet_priority"),
                "far_east_return_required": vessel.get("far_east_return_required"),
                "switch_allowed": voyage.get("route_change_allowed"),
                "switch_approval_owner": None if voyage.get("route_change_allowed") else "FLEET_MANAGEMENT",
                "allocation_status": "ALLOCATED",
                "constraint_summary": vessel.get("availability_note"),
                "data_source": vessel.get("data_source"),
                "last_updated": vessel.get("last_updated"),
            }
        )

    port_compatibility: list[dict[str, Any]] = []
    for vessel in vessels:
        voyage = voyages_by_vessel[vessel["vessel_id"]]
        codes = list(dict.fromkeys(code.strip() for code in str(voyage.get("route_plan") or "").split(">") if code.strip()))
        for code in codes:
            port = ports_by_code.get(code)
            if port is None:
                raise FleetDataError(f"Voyage {voyage['voyage_id']} references unknown port {code}")
            compatible = (
                _require_number(vessel, "design_draft_m", f"Vessel {vessel['vessel_id']}") <= _require_number(port, "max_draft_m", f"Port {code}")
                and _require_number(vessel, "loa_m", f"Vessel {vessel['vessel_id']}") <= _require_number(port, "max_loa_m", f"Port {code}")
            )
            if code == "CNLYG" and (vessel.get("year_built") or 0) < 2014:
                compatible = False
            port_compatibility.append(
                {
                    "compatibility_id": f"PC-{vessel['vessel_id']}-{code}",
                    "vessel_id": vessel["vessel_id"],
                    "port_code": code,
                    "is_compatible": compatible,
                    "restriction_reason": "NONE" if compatible else port.get("port_restriction_note"),
                    "max_allowed_draft_m": port.get("max_draft_m"),
                    "vessel_design_draft_m": vessel.get("design_draft_m"),
                    "max_allowed_loa_m": port.get("max_loa_m"),
                    "vessel_loa_m": vessel.get("loa_m"),
                    "approval_required": not compatible,
                    "valid_from": None,
                    "valid_to": None,
                    "data_source": port.get("data_source"),
                }
            )

    payload: dict[str, Any] = {
        "vessel": vessels,
        "cargo": raw["CARGO"],
        "port": raw["PORT"],
        "voyage": raw["VOYAGE"],
        "vessel_live": raw["VESSEL_LIVE"],
        "maintenance_schedule": maintenance_schedule,
        "charter_status": charter_status,
        "fleet_allocation": fleet_allocation,
        "port_compatibility": port_compatibility,
    }
    content_hash = _content_hash(payload)
    previous_meta = previous_payload.get("meta", {}) if previous_payload else {}
    if previous_meta.get("content_hash") == content_hash:
        effective_generated_at = previous_meta.get("generated_at")
        source_metadata = previous_meta.get("source")
    else:
        effective_generated_at = generated_at or datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")
        source_metadata = {
            "type": "OneDrive public Excel",
            "url": source_url,
            "workbook_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "sheets": {name: len(records) for name, records in raw.items()},
        }

    payload["meta"] = {
        "filename": "fleet-source.xlsx",
        "as_of": max(str(vessel.get("last_updated") or "") for vessel in vessels),
        "timezone": "미지정",
        "classification": "SIMULATED",
        "generated_at": effective_generated_at,
        "content_hash": content_hash,
        "source": source_metadata,
        "counts": {key.upper(): len(payload[key]) for key in DATA_KEYS},
        "checks": [
            {"name": "VESSEL.vessel_id 키 중복·누락", "count": 0, "examples": []},
            {"name": "선박·최신 위치·항차 1:1 연결", "count": 0, "examples": []},
            {"name": "화물 선박·항차 외래키", "count": 0, "examples": []},
            {"name": "잔여·가용 적재량 계산", "count": 0, "examples": []},
            {"name": "3층별 배정 화물·가용량 계산", "count": 0, "examples": []},
            {"name": "항로 항구 코드·접안 조건", "count": 0, "examples": []},
        ],
        "issues": {},
        "original_validation": [
            {"check_name": "Duplicate vessel_id", "issue_count": 0, "status": "PASS", "checked_at": effective_generated_at, "detail": "VESSEL.vessel_id must be unique"},
            {"check_name": "Missing live position", "issue_count": 0, "status": "PASS", "checked_at": effective_generated_at, "detail": "Every vessel must have one VESSEL_LIVE row"},
            {"check_name": "Invalid route port", "issue_count": 0, "status": "PASS", "checked_at": effective_generated_at, "detail": "Every route_plan code must exist in PORT"},
            {"check_name": "Deck capacity mismatch", "issue_count": 0, "status": "PASS", "checked_at": effective_generated_at, "detail": "Deck availability must reconcile to assigned cargo"},
        ],
    }
    validate_payload(payload)
    return payload


def load_json(path: str | Path) -> dict[str, Any] | None:
    candidate = Path(path)
    if not candidate.is_file():
        return None
    with candidate.open("r", encoding="utf-8") as handle:
        value = json.load(handle)
    if not isinstance(value, dict):
        raise FleetDataError(f"Previous JSON must contain an object: {candidate}")
    return value


def write_json_atomic(payload: dict[str, Any], output_path: str | Path) -> None:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(f".{output.name}.tmp")
    with temporary.open("w", encoding="utf-8", newline="\n") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2, sort_keys=False)
        handle.write("\n")
    temporary.replace(output)


def write_javascript_atomic(payload: dict[str, Any], output_path: str | Path) -> None:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(f".{output.name}.tmp")
    with temporary.open("w", encoding="utf-8", newline="\n") as handle:
        handle.write("window.FLEET_DATA=")
        json.dump(payload, handle, ensure_ascii=False, separators=(",", ":"))
        handle.write(";\n")
    temporary.replace(output)
